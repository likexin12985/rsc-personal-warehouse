"""Build both public catalogues from exact reviewed XLSX cells; never fetch or deploy.

The source workbook stays outside the repository. Each normalized row must keep
its exact sheet and row number so the selected public fields can be reviewed.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import re
import tempfile
from zipfile import BadZipFile, ZipFile

ROOT = Path(__file__).resolve().parents[1]
SOURCE_URL = "https://wbenergy.feishu.cn/sheets/ZVnis9SUzhfiU9t9EK4csLvEnNg?sheet=1eOLBU"
OUTPUTS = (ROOT / "frontend/src/knowledge-catalog.json", ROOT / "miniprogram/data/knowledge-catalog.json")
TEXT_LIMITS = {"code": 80, "name": 500, "category": 100, "model": 1000, "note": 3000, "sourceSheet": 100}
MAX_SOURCE_BYTES = 50_000_000
MAX_XML_BYTES = 100_000_000
MAX_CELLS = 2_000_000


def _exact(value, keys, label):
    if not isinstance(value, dict) or set(value) != set(keys):
        raise ValueError(f"{label}: unexpected or missing fields")
    return value


def _integer(value, low, high, label):
    if type(value) is not int or not low <= value <= high:
        raise ValueError(f"{label}: integer outside supported range")
    return value


def _reason(value):
    if not isinstance(value, str) or not value.strip() or len(value) > 500:
        raise ValueError("every exclusion requires a review reason")


def _json_without_duplicates(raw):
    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                raise ValueError("duplicate review JSON key")
            result[key] = value
        return result
    return json.loads(raw, object_pairs_hook=pairs)


def _workbook(content):
    if not 0 < len(content) <= MAX_SOURCE_BYTES:
        raise ValueError("source workbook is empty or exceeds the import limit")
    try:
        from defusedxml.ElementTree import iterparse
        from defusedxml.common import DefusedXmlException
        from xml.etree.ElementTree import ParseError
        from openpyxl.utils.cell import coordinate_to_tuple
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            names = [entry.filename for entry in members]
            if (len(names) != len(set(names)) or len(names) > 10000
                    or sum(entry.file_size for entry in members) > MAX_XML_BYTES
                    or any(entry.flag_bits & 1 for entry in members)
                    or any(name.startswith('/') or '..' in name.split('/') for name in names)
                    or any('vbaproject' in name.lower() or name.startswith('xl/externalLinks/') for name in names)):
                raise ValueError("source archive is unsupported or exceeds safe import bounds")
            if 'xl/workbook.xml' not in names or '[Content_Types].xml' not in names:
                raise ValueError("source must be a complete XLSX workbook")
            # Bound actual coordinates before openpyxl can allocate padded rows.
            # Cached worksheet dimensions are deliberately ignored.
            cells = 0
            for name in names:
                if not re.fullmatch(r'xl/worksheets/[^/]+\.xml', name):
                    continue
                max_row = max_column = current_row = previous_row = previous_column = 0
                with archive.open(name) as source:
                    for event, element in iterparse(source, events=('start', 'end')):
                        tag = element.tag.rsplit('}', 1)[-1]
                        if event == 'start':
                            if tag != 'row':
                                continue
                            raw_row = element.attrib.get('r', '')
                            if not re.fullmatch(r'[1-9][0-9]{0,5}', raw_row):
                                raise ValueError("source row coordinate is invalid")
                            row = int(raw_row)
                            if not previous_row < row <= 100000:
                                raise ValueError("source row exceeds the import limit")
                            current_row = previous_row = row
                            previous_column = 0
                            max_row = max(max_row, row)
                            continue
                        if tag == 'c':
                            coordinate = element.attrib.get('r', '')
                            if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,5}', coordinate):
                                raise ValueError("source cell coordinate is invalid")
                            row, column = coordinate_to_tuple(coordinate)
                            if row != current_row or not previous_column < column <= 256:
                                raise ValueError("source dimensions exceed the import limit")
                            previous_column = column
                            max_row, max_column = max(max_row, row), max(max_column, column)
                        elif tag == 'row':
                            current_row = 0
                        element.clear()
                cells += max_row * max_column
                if cells > MAX_CELLS:
                    raise ValueError("source workbook exceeds the bounded cell scan")
        from openpyxl import load_workbook
        from openpyxl.xml import DEFUSEDXML
        if not DEFUSEDXML:
            raise ValueError("protected XML parsing is required; install requirements-public-knowledge.txt")
        try:
            return load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except (ValueError, TypeError, IndexError):
            raise ValueError("source workbook metadata is invalid or unsupported") from None
    except ImportError:
        raise ValueError("install scripts/requirements-public-knowledge.txt before importing") from None
    except (BadZipFile, KeyError, OSError):
        raise ValueError("source must be a readable complete XLSX workbook") from None
    except (DefusedXmlException, ParseError):
        raise ValueError("source workbook XML is invalid or unsafe") from None


def _field_binding(sheet, descriptor, first_row, label):
    from openpyxl.utils.cell import column_index_from_string
    _exact(descriptor, ('column', 'headerCell', 'headerText'), label)
    column, header = descriptor['column'], descriptor['headerCell']
    if not isinstance(column, str) or not re.fullmatch('[A-Z]{1,3}', column):
        raise ValueError(f"{label}: invalid source column")
    match = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,5})', header) if isinstance(header, str) else None
    if not match or match[1] != column or int(match[2]) >= first_row:
        raise ValueError(f"{label}: header must be in the same column above the selected data")
    expected = descriptor['headerText']
    if not isinstance(expected, str) or not expected.strip() or len(expected) > 500:
        raise ValueError(f"{label}: exact source header text is required")
    index = column_index_from_string(column)
    if index > 256:
        raise ValueError(f"{label}: source column exceeds the import limit")
    cell = sheet[header]
    if cell.data_type == 'f' or cell.value != expected:
        raise ValueError(f"{label}: source header changed or does not match the review")
    return index


def extract_reviewed_rows(content: bytes, review: object) -> tuple[list, dict]:
    """Read only selected source text. No invented models, categories or notes."""
    _exact(review, ('schemaVersion', 'sourceUrl', 'sourceSha256', 'sheets'), 'review')
    digest = hashlib.sha256(content).hexdigest()
    if (type(review['schemaVersion']) is not int or review['schemaVersion'] != 1
            or review['sourceUrl'] != SOURCE_URL or review['sourceSha256'] != digest):
        raise ValueError("review must bind the selected source and exact current export digest")
    plans = review['sheets']
    if not isinstance(plans, list) or not plans or len(plans) > 100:
        raise ValueError("review must explicitly include or exclude every worksheet")
    names = [plan.get('name') if isinstance(plan, dict) else None for plan in plans]
    if any(not isinstance(name, str) or not name for name in names) or len(set(names)) != len(names):
        raise ValueError("worksheet review names must be exact and unique")
    workbook = _workbook(content)
    try:
        if set(names) != set(workbook.sheetnames):
            raise ValueError("worksheet set changed or review does not cover every worksheet")
        rows, summary, total_cells = [], [], 0
        for number, plan in enumerate(plans, 1):
            label = f"worksheet {number}"
            if plan.get('action') == 'exclude':
                _exact(plan, ('name', 'action', 'reason'), label)
                _reason(plan['reason'])
                summary.append(dict(worksheet=number, action='excluded'))
                continue
            _exact(plan, ('name', 'action', 'includeHidden', 'firstRow', 'lastRow', 'expectedRecords', 'fields', 'excludedRows'), label)
            if plan['action'] != 'include' or type(plan['includeHidden']) is not bool:
                raise ValueError(f"{label}: explicit inclusion is required")
            sheet = workbook[plan['name']]
            if sheet.sheet_state != 'visible' and not plan['includeHidden']:
                raise ValueError(f"{label}: hidden worksheet was not explicitly reviewed for inclusion")
            first = _integer(plan['firstRow'], 2, 100000, label)
            last = _integer(plan['lastRow'], first, 100000, label)
            expected = _integer(plan['expectedRecords'], 1, 20000, label)
            # Do not trust OOXML's cached dimensions: a stale A1:A1 can hide rows.
            sheet.reset_dimensions()
            bounds = sheet.calculate_dimension(force=True)
            if sheet.max_row > 100000 or sheet.max_column > 256:
                raise ValueError(f"{label}: source dimensions exceed the import limit")
            total_cells += sheet.max_row * sheet.max_column
            if total_cells > MAX_CELLS:
                raise ValueError("source workbook exceeds the bounded cell scan")
            fields = _exact(plan['fields'], ('code', 'name', 'category', 'model', 'note'), label)
            bindings = {}
            for field, descriptor in fields.items():
                if field in ('model', 'note') and descriptor is None:
                    bindings[field] = None
                elif field == 'category' and descriptor == {'sheetName': True} and type(descriptor['sheetName']) is bool:
                    bindings[field] = 'sheet'
                else:
                    bindings[field] = _field_binding(sheet, descriptor, first, f"{label} {field}")
            indices = [value for value in bindings.values() if type(value) is int]
            if len(indices) != len(set(indices)):
                raise ValueError(f"{label}: public fields cannot silently reuse the same source column")
            excluded = plan['excludedRows']
            if not isinstance(excluded, list):
                raise ValueError(f"{label}: excluded rows must be an explicit list")
            skipped = set()
            for entry in excluded:
                _exact(entry, ('row', 'reason'), label)
                index = _integer(entry['row'], first, last, label); _reason(entry['reason'])
                if index in skipped:
                    raise ValueError(f"{label}: duplicate excluded row")
                skipped.add(index)
            count, blanks, final_content_row = 0, 0, 0
            for row_number, cells in enumerate(sheet.iter_rows(), 1):
                nonempty = any(cell.value is not None and cell.value != '' for cell in cells)
                if nonempty:
                    final_content_row = row_number
                if row_number < first or row_number > last or row_number in skipped:
                    continue
                if not nonempty:
                    blanks += 1
                    continue
                item = {'sourceSheet': sheet.title, 'sourceRow': row_number}
                for field, binding in bindings.items():
                    if binding is None:
                        value = ''
                    elif binding == 'sheet':
                        value = sheet.title
                    else:
                        cell = cells[binding - 1] if binding <= len(cells) else None
                        value = cell.value if cell is not None else None
                        if cell is not None and cell.data_type in ('f', 'e'):
                            raise ValueError(f"{label} row {row_number} {field}: formula or error is not a reviewed literal")
                        if value is None and field in ('model', 'note'):
                            value = ''
                        if not isinstance(value, str):
                            raise ValueError(f"{label} row {row_number} {field}: source text required; no numeric or date coercion")
                    item[field] = value.strip()
                rows.append(item); count += 1
                if len(rows) > 20000:
                    raise ValueError("reviewed records exceed the catalog limit")
            if final_content_row > last or last > sheet.max_row:
                raise ValueError(f"{label}: reviewed row range omits source content or exceeds the source")
            if count != expected:
                raise ValueError(f"{label}: extracted record count differs from the reviewed count")
            summary.append(dict(worksheet=number, action='included', records=count, excludedRows=len(skipped), blankRows=blanks, dimensions=bounds))
        return rows, dict(sourceSha256=digest, sheets=summary, records=len(rows))
    finally:
        workbook.close()


def replace_catalogs(payload: str) -> None:
    """Stage both outputs before replacement; restore owned writes on an I/O failure.

    This is not a cross-file filesystem transaction. Existing build/release
    checks must still reject mismatched outputs after process/power loss.
    """
    before = {output: output.read_bytes() for output in OUTPUTS}
    temporary_paths, written = {}, []
    try:
        for output in OUTPUTS:
            with tempfile.NamedTemporaryFile(mode='wb', dir=output.parent, delete=False) as temporary:
                temporary_paths[output] = Path(temporary.name)
                temporary.write(payload.encode('utf-8')); temporary.flush(); os.fsync(temporary.fileno())
        if any(output.read_bytes() != original for output, original in before.items()):
            raise ValueError("catalog changed during import; preserve it and retry after review")
        for output in OUTPUTS:
            if output.read_bytes() != before[output]:
                raise ValueError("catalog changed during import; preserve it and retry after review")
            temporary_paths[output].replace(output); written.append(output)
        if any(output.read_text(encoding='utf-8') != payload for output in OUTPUTS):
            raise ValueError("catalog readback differs from the reviewed source")
    except (OSError, ValueError):
        for output in written:
            if output.read_bytes() == payload.encode('utf-8'):
                output.write_bytes(before[output])
        raise
    finally:
        for temporary in temporary_paths.values():
            temporary.unlink(missing_ok=True)


def build_catalog(rows: object, *, source_url: str, verified_at: str, source_sha256: str) -> dict:
    if source_url != SOURCE_URL:
        raise ValueError("source URL must match the explicitly selected Feishu sheet")
    timestamp = datetime.fromisoformat(verified_at.replace("Z", "+00:00"))
    if timestamp.tzinfo is None or timestamp > datetime.now(timezone.utc):
        raise ValueError("verified-at requires a timezone and cannot be in the future")
    if not re.fullmatch(r"[0-9a-f]{64}", source_sha256):
        raise ValueError("source export requires a SHA256 digest")
    if not isinstance(rows, list) or not 1 <= len(rows) <= 20000:
        raise ValueError("reviewed rows must contain between 1 and 20000 records")
    items = []
    seen = set()
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, dict) or set(row) != set(TEXT_LIMITS) | {"sourceRow"}:
            raise ValueError(f"record {index}: only the documented public fields are allowed")
        item = {}
        for field, limit in TEXT_LIMITS.items():
            value = row[field]
            if not isinstance(value, str) or len(value) > limit or any(ord(c) < 32 and c not in "\n\t" for c in value):
                raise ValueError(f"record {index}: invalid {field}")
            item[field] = value if field == 'sourceSheet' else value.strip()
        if any(not item[key].strip() for key in ("code", "name", "category", "sourceSheet")):
            raise ValueError(f"record {index}: code/name/category/sourceSheet cannot be empty")
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._/\-]{0,79}", item["code"]):
            raise ValueError(f"record {index}: invalid material code")
        if type(row["sourceRow"]) is not int or not 1 <= row["sourceRow"] <= 1000000:
            raise ValueError(f"record {index}: sourceRow must be a positive worksheet row number")
        item["sourceRow"] = row["sourceRow"]
        key = (item["sourceSheet"], item["sourceRow"], item["code"])
        if key in seen:
            raise ValueError(f"record {index}: duplicate source row and material code")
        seen.add(key)
        items.append(item)
    return {
        "schemaVersion": 1, "status": "ready", "sourceUrl": source_url,
        "verifiedAt": timestamp.isoformat(), "sourceSha256": source_sha256,
        "items": sorted(items, key=lambda row: (row["code"], row["sourceSheet"], row["sourceRow"])),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--review-plan", type=Path, required=True, help="reviewed exact XLSX worksheet/column/row bindings")
    parser.add_argument("--source-export", type=Path, required=True, help="original exported workbook, kept outside Git")
    parser.add_argument("--source-url", required=True)
    parser.add_argument("--verified-at", required=True, help="actual source verification time in ISO8601 with timezone")
    parser.add_argument("--write", action="store_true", help="replace both local catalogues after validation")
    args = parser.parse_args()
    try:
        if not 0 < args.review_plan.stat().st_size <= 10_000_000 or not 0 < args.source_export.stat().st_size <= MAX_SOURCE_BYTES:
            raise ValueError("input file is empty or exceeds the import limit")
        if args.source_export.suffix.lower() != '.xlsx' or args.source_export.resolve().is_relative_to(ROOT.parent):
            raise ValueError("keep the original XLSX export outside the Git repository")
        content = args.source_export.read_bytes()
        review_content = args.review_plan.read_bytes()
        rows, evidence = extract_reviewed_rows(content, _json_without_duplicates(review_content))
        catalog = build_catalog(rows, source_url=args.source_url, verified_at=args.verified_at, source_sha256=evidence['sourceSha256'])
        payload = json.dumps(catalog, ensure_ascii=False, indent=2) + "\n"
        if args.write:
            replace_catalogs(payload)
        print(json.dumps({"status": "written" if args.write else "validated", **evidence,
            "reviewSha256": hashlib.sha256(review_content).hexdigest(), "catalogSha256": hashlib.sha256(payload.encode()).hexdigest()}))
    except (ValueError, OSError) as error:
        parser.exit(2, f"Import rejected: {error}\n")


if __name__ == "__main__":
    main()
