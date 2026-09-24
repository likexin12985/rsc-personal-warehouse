from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock
from uuid import uuid4

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from app.database import get_db
from app.dependencies import get_formal_principal
from app.config import get_settings
from app.formal_services.opening_count_import_source import OpeningCountImportSourceError
from app.formal_services.opening_count_import_workbook import ImportRowError
from app.formal_services.opening_count_import_workbook import HEADERS, SHEET_NAME, MAX_FILE_BYTES
from app.routers import formal_opening_stocktake
from app.routers.formal_files import get_formal_file_storage_adapter


MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BASE = "/api/v1/stocktakes/opening/imports/opening-count"


@pytest.fixture()
def client_with_scope():
    db = SimpleNamespace(commit=Mock(), rollback=Mock())
    principal = SimpleNamespace(allows=lambda _db, resource, action, **_kw:
                                (resource, action) == ("stocktake", "count"))
    api = FastAPI()
    api.include_router(formal_opening_stocktake.router, prefix="/api")
    api.dependency_overrides[get_db] = lambda: db
    api.dependency_overrides[get_formal_principal] = lambda: principal
    with TestClient(api) as client:
        yield client, db, principal


def _workbook(*rows):
    book = Workbook()
    book.active.title = SHEET_NAME
    book.active.append(HEADERS)
    for row in rows:
        book.active.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _row(quantity="1.000", material="ADQCPN0123"):
    return material, "sku_code", "new", "available", quantity, None, None, None, None, ""


def test_template_and_format_preview_keep_stock_untouched(client_with_scope):
    client, db, _ = client_with_scope
    template = client.get(BASE + "/template")
    assert template.status_code == 200 and template.headers["cache-control"] == "no-store"
    book = load_workbook(BytesIO(template.content), read_only=True)
    try:
        assert book.sheetnames == [SHEET_NAME]
        assert next(book.active.values) == HEADERS
    finally:
        book.close()
    response = client.post(BASE + "/format-check", content=_workbook(_row()), headers={"Content-Type": MIME})
    assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
    body = response.json()
    assert body["format_valid"] and body["row_count"] == 1 and body["errors"] == []
    assert len(body["source_sha256"]) == len(body["payload_sha256"]) == 64
    assert "observations" not in body and "ADQCPN0123" not in response.text
    db.commit.assert_not_called()
    db.rollback.assert_not_called()


def test_error_report_returns_only_controlled_error_metadata(client_with_scope):
    client, db, _ = client_with_scope
    source = _workbook(_row(quantity="0", material="PRIVATE-MATERIAL-CODE"))
    preview = client.post(BASE + "/format-check", content=source, headers={"Content-Type": MIME})
    assert preview.status_code == 200
    assert not preview.json()["format_valid"]
    assert preview.json()["payload_sha256"] is None
    report = client.post(BASE + "/error-report", content=source, headers={"Content-Type": MIME})
    assert report.status_code == 200 and report.headers["content-type"].startswith(MIME)
    assert report.headers["cache-control"] == "no-store"
    book = load_workbook(BytesIO(report.content), read_only=True)
    try:
        entries = list(book.active.values)
        assert entries[1][1:4] == (2, "counted_qty", "quantity_invalid")
        assert "PRIVATE-MATERIAL-CODE" not in str(entries)
    finally:
        book.close()
    db.commit.assert_not_called()


def test_unauthorized_and_invalid_transport_fail_closed(client_with_scope):
    client, db, principal = client_with_scope
    principal.allows = lambda *_args, **_kw: False
    response = client.get(BASE + "/template")
    assert response.status_code == 403
    principal.allows = lambda *_args, **_kw: True
    assert client.post(BASE + "/format-check", content=b"bad", headers={"Content-Type": "text/plain"}).status_code == 415
    assert client.post(BASE + "/format-check", content=b"bad", headers={"Content-Type": MIME}).status_code == 422
    oversized = client.post(BASE + "/format-check", content=b"x" * (MAX_FILE_BYTES + 1), headers={"Content-Type": MIME})
    assert oversized.status_code == 413
    assert client.post(BASE + "/error-report", content=_workbook(_row()), headers={"Content-Type": MIME}).status_code == 409
    db.commit.assert_not_called()


def test_private_business_preview_route_binds_ids_and_returns_only_evidence(monkeypatch):
    captured = []
    db = SimpleNamespace()
    principal = SimpleNamespace(allows=lambda *_args, **_kwargs: True)
    storage = SimpleNamespace(provider_code="aliyun_oss_v2")
    api = FastAPI()
    api.include_router(formal_opening_stocktake.router, prefix="/api")
    api.dependency_overrides[get_db] = lambda: db
    api.dependency_overrides[get_formal_principal] = lambda: principal
    api.dependency_overrides[get_settings] = lambda: SimpleNamespace(
        file_storage_configuration_ready=lambda: True,
    )
    api.dependency_overrides[get_formal_file_storage_adapter] = lambda: storage
    api.dependency_overrides[formal_opening_stocktake.get_opening_import_session_factory] = lambda: object()

    def preview(*args, **kwargs):
        captured.append(kwargs)
        return SimpleNamespace(
            source_sha256="a" * 64, payload_sha256="b" * 64,
            row_count=1, ready=False,
            count=SimpleNamespace(
                task_version=7, actor_authorization_version=4,
                request_sha256="c" * 64,
            ),
            errors=(ImportRowError(
                row=4, field="row", code="pending_verification",
                message="物料尚未核实",
            ),),
        )

    monkeypatch.setattr(
        formal_opening_stocktake.import_prevalidation,
        "prevalidate_authorized_opening_count_import", preview,
    )
    ids = {name: str(uuid4()) for name in (
        "source_file_id", "task_id", "round_id", "scope_id",
    )}
    headers = {"Idempotency-Key": "opening-import-123456", "X-Request-ID": "req-123456"}
    with TestClient(api) as client:
        response = client.post(BASE + "/business-check", json=ids, headers=headers)
        assert response.status_code == 200
        assert response.headers["cache-control"] == "no-store"
        body = response.json()
        assert body["ready"] is False and body["task_version"] == 7
        assert body["errors"] == [{
            "row": 4, "field": "row", "code": "pending_verification",
            "message": "物料尚未核实",
        }]
        assert str(captured[0]["file_id"]) == ids["source_file_id"]
        assert captured[0]["idempotency_key"] == headers["Idempotency-Key"]
        assert "observations" not in response.text
        assert client.post(BASE + "/business-check", json={**ids, "scope_id": str(uuid4()), "extra": True}, headers=headers).status_code == 422

        def forbidden(*args, **kwargs):
            raise OpeningCountImportSourceError("opening_import_source_unavailable", 404, "源文件不可用")

        monkeypatch.setattr(
            formal_opening_stocktake.import_prevalidation,
            "prevalidate_authorized_opening_count_import", forbidden,
        )
        rejected = client.post(BASE + "/business-check", json=ids, headers=headers)
        assert rejected.status_code == 404 and rejected.headers["cache-control"] == "no-store"
        assert rejected.json()["detail"]["code"] == "opening_import_source_unavailable"
        api.dependency_overrides[get_formal_file_storage_adapter] = lambda: None
        disabled = client.post(BASE + "/business-check", json=ids, headers=headers)
        assert disabled.status_code == 503
        assert disabled.json()["detail"]["code"] == "opening_import_storage_disabled"
    assert len(captured) == 1
