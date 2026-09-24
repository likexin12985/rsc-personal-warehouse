"""Synthetic XLSX only: prove each published field comes from reviewed cells."""
from copy import deepcopy
from datetime import datetime
import hashlib
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

from openpyxl import Workbook
import pytest

from test_public_knowledge_catalog import module


def workbook_bytes(edit=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "测试分类"
    sheet.append(["物料编码", "备件名称", "型号", "公开说明", "私有字段"])
    sheet.append(["001ABC", "合成备件", None, "仅测试夹具", "PRIVATE_SENTINEL"])
    sheet.append(["TEST002", "合成备件二", "型号二", None, "PRIVATE_SENTINEL"])
    book.create_sheet("不公开页").append(["PRIVATE_SENTINEL"])
    if edit:
        edit(book)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def binding(column, header):
    return dict(column=column, headerCell=column + '1', headerText=header)


def review(content):
    return dict(schemaVersion=1, sourceUrl=module.SOURCE_URL,
        sourceSha256=hashlib.sha256(content).hexdigest(), sheets=[dict(
            name="测试分类", action="include", includeHidden=False,
            firstRow=2, lastRow=3, expectedRecords=2,
            fields=dict(code=binding('A', '物料编码'), name=binding('B', '备件名称'),
                category={'sheetName': True}, model=binding('C', '型号'), note=binding('D', '公开说明')),
            excludedRows=[]), dict(name="不公开页", action="exclude", reason="未选定公开字段")])


def rewrite_archive(content, transform):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, 'w') as target:
        for entry in source.infolist():
            target.writestr(entry.filename, transform(entry.filename, source.read(entry)))
    return output.getvalue()


def test_exact_cells_preserve_leading_zero_unknown_model_and_source_coordinates():
    content = workbook_bytes()
    rows, evidence = module.extract_reviewed_rows(content, review(content))
    assert rows[0] == dict(code='001ABC', name='合成备件', category='测试分类', model='',
        note='仅测试夹具', sourceSheet='测试分类', sourceRow=2)
    assert rows[1]['sourceRow'] == 3 and rows[1]['model'] == '型号二'
    assert evidence['records'] == 2
    assert evidence['sheets'] == [dict(worksheet=1, action='included', records=2,
        excludedRows=0, blankRows=0, dimensions='A1:E3'), dict(worksheet=2, action='excluded')]
    assert 'PRIVATE_SENTINEL' not in json.dumps([rows, evidence])


@pytest.mark.parametrize('change', [
    lambda p: p.update(sourceSha256='0' * 64),
    lambda p: p.update(sourceUrl='https://example.com'),
    lambda p: p.update(schemaVersion=True),
    lambda p: p['sheets'].pop(),
    lambda p: p['sheets'].append(deepcopy(p['sheets'][0])),
    lambda p: p['sheets'][1].update(reason=''),
    lambda p: p['sheets'][0].update(lastRow=2),
    lambda p: p['sheets'][0].update(lastRow=4),
    lambda p: p['sheets'][0].update(expectedRecords=1),
    lambda p: p['sheets'][0].update(firstRow=True),
    lambda p: p['sheets'][0]['fields']['code'].update(headerText='changed'),
    lambda p: p['sheets'][0]['fields']['code'].update(headerCell='B1'),
    lambda p: p['sheets'][0]['fields']['code'].update(headerCell='A2'),
    lambda p: p['sheets'][0]['fields'].update(name=binding('A', '物料编码')),
    lambda p: p['sheets'][0]['fields'].update(category={'sheetName': 1}),
    lambda p: p['sheets'][0]['fields'].update(privateField=binding('E', '私有字段')),
    lambda p: p['sheets'][0].update(excludedRows=[dict(row=2, reason='')]),
    lambda p: p['sheets'][0].update(excludedRows=[dict(row=4, reason='范围外')]),
    lambda p: p['sheets'][0].update(excludedRows=[dict(row=2, reason='重复')] * 2),
])
def test_review_drift_missing_coverage_or_ambiguous_mapping_is_rejected(change):
    content = workbook_bytes(); plan = review(content); change(plan)
    with pytest.raises(ValueError):
        module.extract_reviewed_rows(content, plan)


def test_explicit_row_exclusion_and_unselected_model_do_not_infer_values():
    content = workbook_bytes(); plan = review(content)
    plan['sheets'][0].update(excludedRows=[dict(row=2, reason='不在公开范围')], expectedRecords=1)
    plan['sheets'][0]['fields']['model'] = None
    rows, evidence = module.extract_reviewed_rows(content, plan)
    assert len(rows) == 1 and rows[0]['sourceRow'] == 3 and rows[0]['model'] == ''
    assert evidence['sheets'][0]['excludedRows'] == 1


def test_exact_sheet_title_is_not_trimmed_into_a_different_source_coordinate():
    content = workbook_bytes(lambda b: setattr(b.active, 'title', ' 测试分类 '))
    plan = review(content); plan['sheets'][0]['name'] = ' 测试分类 '
    rows, _ = module.extract_reviewed_rows(content, plan)
    catalog = module.build_catalog(rows, source_url=module.SOURCE_URL,
        verified_at='2026-01-01T00:00:00Z', source_sha256=plan['sourceSha256'])
    assert all(row['sourceSheet'] == ' 测试分类 ' for row in catalog['items'])


@pytest.mark.parametrize('value', [123, True, datetime(2026, 1, 1), '=1+1', '#REF!'])
def test_formula_error_numeric_and_date_cells_are_not_public_text(value):
    content = workbook_bytes(lambda b: setattr(b.active['A2'], 'value', value))
    with pytest.raises(ValueError, match='reviewed literal|source text'):
        module.extract_reviewed_rows(content, review(content))


def test_hidden_sheet_requires_exact_review_and_new_sheet_invalidates_coverage():
    content = workbook_bytes(lambda b: setattr(b.active, 'sheet_state', 'hidden'))
    plan = review(content)
    with pytest.raises(ValueError, match='hidden worksheet'):
        module.extract_reviewed_rows(content, plan)
    plan['sheets'][0]['includeHidden'] = True
    assert len(module.extract_reviewed_rows(content, plan)[0]) == 2
    content = workbook_bytes(lambda b: b.create_sheet('新增未审核页'))
    with pytest.raises(ValueError, match='worksheet set'):
        module.extract_reviewed_rows(content, review(content))


def test_stale_cached_dimensions_do_not_hide_tail_records():
    content = rewrite_archive(workbook_bytes(), lambda name, data:
        data.replace(b'<dimension ref="A1:E3"', b'<dimension ref="A1:A1"')
        if name == 'xl/worksheets/sheet1.xml' else data)
    assert len(module.extract_reviewed_rows(content, review(content))[0]) == 2


@pytest.mark.parametrize('old,new', [
    (b'r="A2"', b'r="A999999"'), (b'r="A2"', b'r="XFD2"'),
    (b'r="B2"', b'r="A2"'), (b'<row r="2"', b'<row r="999999"'),
    (b'<row r="2"', b'<row r="1"'), (b'<row r="2"', b'<row r="PRIVATE_SENTINEL"'),
])
def test_invalid_or_unbounded_real_coordinates_rejected_before_padding(old, new):
    content = rewrite_archive(workbook_bytes(), lambda name, data:
        data.replace(old, new) if name == 'xl/worksheets/sheet1.xml' else data)
    with pytest.raises(ValueError) as error:
        module.extract_reviewed_rows(content, review(content))
    assert 'PRIVATE_SENTINEL' not in str(error.value)


def test_xml_entity_and_archive_expansion_are_rejected(monkeypatch):
    content = rewrite_archive(workbook_bytes(), lambda name, data:
        b'<!DOCTYPE worksheet [<!ENTITY x "PRIVATE_SENTINEL">]>' + data
        if name == 'xl/worksheets/sheet1.xml' else data)
    with pytest.raises(ValueError, match='XML is invalid or unsafe'):
        module.extract_reviewed_rows(content, review(content))
    monkeypatch.setattr(module, 'MAX_XML_BYTES', 10)
    content = workbook_bytes()
    with pytest.raises(ValueError, match='safe import bounds'):
        module.extract_reviewed_rows(content, review(content))


@pytest.mark.parametrize('name', ['xl/vbaProject.bin', 'xl/externalLinks/externalLink1.xml', '../private'])
def test_unsupported_archive_members_are_rejected(name):
    stream = BytesIO(workbook_bytes())
    with ZipFile(stream, 'a') as archive:
        archive.writestr(name, 'PRIVATE_SENTINEL')
    content = stream.getvalue()
    with pytest.raises(ValueError, match='safe import bounds'):
        module.extract_reviewed_rows(content, review(content))


def test_duplicate_json_keys_cannot_change_review_silently():
    with pytest.raises(ValueError, match='duplicate review JSON key'):
        module._json_without_duplicates('{"sheets":[],"sheets":[]}')


def output_paths(tmp_path, monkeypatch):
    outputs = (tmp_path / 'web.json', tmp_path / 'mini.json')
    for output in outputs:
        output.write_text('before')
    monkeypatch.setattr(module, 'OUTPUTS', outputs)
    return outputs


def test_second_output_failure_restores_first_and_cleans_staged_files(tmp_path, monkeypatch):
    outputs = output_paths(tmp_path, monkeypatch)
    replace = Path.replace
    def fail_second(path, target):
        if target == outputs[1]:
            raise OSError('synthetic disk failure')
        return replace(path, target)
    monkeypatch.setattr(Path, 'replace', fail_second)
    with pytest.raises(OSError, match='disk failure'):
        module.replace_catalogs('after')
    assert [p.read_text() for p in outputs] == ['before', 'before']
    assert set(tmp_path.iterdir()) == set(outputs)


def test_concurrent_edit_to_second_output_is_preserved(tmp_path, monkeypatch):
    outputs = output_paths(tmp_path, monkeypatch)
    replace = Path.replace
    def edit_second(path, target):
        result = replace(path, target)
        if target == outputs[0]:
            outputs[1].write_text('concurrent-edit')
        return result
    monkeypatch.setattr(Path, 'replace', edit_second)
    with pytest.raises(ValueError, match='changed during import'):
        module.replace_catalogs('after')
    assert [p.read_text() for p in outputs] == ['before', 'concurrent-edit']
    assert set(tmp_path.iterdir()) == set(outputs)


def test_stage_failure_leaves_both_catalogs_untouched(tmp_path, monkeypatch):
    outputs = output_paths(tmp_path, monkeypatch)
    def failure(_fd):
        raise OSError('synthetic sync failure')
    monkeypatch.setattr(module.os, 'fsync', failure)
    with pytest.raises(OSError, match='sync failure'):
        module.replace_catalogs('after')
    assert [p.read_text() for p in outputs] == ['before', 'before']
    assert set(tmp_path.iterdir()) == set(outputs)


def test_real_cli_dry_run_then_two_output_readback_and_source_preservation(tmp_path):
    root = tmp_path / 'isolated-repo' / 'cloud_oam'
    script = root / 'scripts/import_public_knowledge.py'
    script.parent.mkdir(parents=True)
    script.write_bytes(Path(module.__file__).read_bytes())
    outputs = [root / p.relative_to(module.ROOT) for p in module.OUTPUTS]
    for output in outputs:
        output.parent.mkdir(parents=True); output.write_text('before')
    content = workbook_bytes()
    source = tmp_path / 'source.xlsx'; source.write_bytes(content)
    plan = tmp_path / 'review.json'; plan.write_text(json.dumps(review(content)))
    args = [sys.executable, str(script), '--review-plan', str(plan), '--source-export', str(source),
        '--source-url', module.SOURCE_URL, '--verified-at', '2026-01-01T00:00:00Z']
    result = subprocess.run(args, capture_output=True, text=True, check=False)
    assert result.returncode == 0, result.stderr
    evidence = json.loads(result.stdout)
    assert evidence['status'] == 'validated' and evidence['records'] == 2
    assert [p.read_text() for p in outputs] == ['before', 'before']
    result = subprocess.run(args + ['--write'], capture_output=True, text=True, check=False)
    assert result.returncode == 0, result.stderr
    evidence = json.loads(result.stdout)
    assert evidence['status'] == 'written'
    assert outputs[0].read_bytes() == outputs[1].read_bytes()
    assert hashlib.sha256(outputs[0].read_bytes()).hexdigest() == evidence['catalogSha256']
    assert source.read_bytes() == content
    assert 'PRIVATE_SENTINEL' not in outputs[0].read_text() + result.stdout + result.stderr
    # Old arbitrary normalized rows cannot be used as source provenance.
    result = subprocess.run(args + ['--rows', str(plan), '--write'], capture_output=True, text=True, check=False)
    assert result.returncode == 2 and 'unrecognized arguments' in result.stderr
    assert hashlib.sha256(outputs[0].read_bytes()).hexdigest() == evidence['catalogSha256']
    # Keep raw source out of Git, including paths resolved through a symlink.
    internal_source = root / 'source.xlsx'; internal_source.write_bytes(content)
    source.unlink(); source.symlink_to(internal_source)
    result = subprocess.run(args + ['--write'], capture_output=True, text=True, check=False)
    assert result.returncode == 2 and 'outside the Git repository' in result.stderr
    assert hashlib.sha256(outputs[0].read_bytes()).hexdigest() == evidence['catalogSha256']
