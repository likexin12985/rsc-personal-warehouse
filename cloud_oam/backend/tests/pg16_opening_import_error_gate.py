"""Real API-role error publication with synthetic private-object transport."""

from copy import deepcopy
from dataclasses import replace
from hashlib import md5, sha256
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import event, select, text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.orm import Session

from formal_file_integrity import FileUploadIntentInput, StoredObjectHead, _prepare_upload
from app.formal_access import load_formal_principal
from app.foundation_models import AuditEvent, FileJob, FileObject
from app.formal_services.opening_count_import_error_jobs import (
    prepare_persisted_opening_count_error, publish_persisted_opening_count_error,
)
from app.formal_services.opening_count_import_error_report import MIME, write_or_recover_opening_count_error
from app.formal_services.opening_count_import_jobs import (
    OpeningCountImportJobError, confirm_persisted_opening_count_import,
    read_opening_count_import_job,
)
from app.formal_services.opening_count_import_source import AuthorizedOpeningCountSource
from app.formal_services.opening_count_import_workbook import FIELDS
from app.formal_services.opening_stocktake_count import _storage_hash
from pg16_opening_source_purpose_gate import _file


def exercise_import_errors(api, *, actor_id, other_actor_id, data, defaults):
    book = load_workbook(BytesIO(data))
    book.active.cell(2, FIELDS.index("counted_qty") + 1, "not-a-quantity")
    stream = BytesIO()
    book.save(stream)
    book.close()
    payload = stream.getvalue()
    digest = sha256(payload).hexdigest()
    prepared = _prepare_upload(FileUploadIntentInput(purpose="opening_count_import",
        original_filename="期初盘点.xlsx", size_bytes=len(payload), mime_type=MIME, sha256=digest),
        maximum_size_bytes=8 * 1024 * 1024)
    with Session(api, expire_on_commit=False) as db:
        actor = load_formal_principal(db, actor_id)
        source = _file(user_id=actor_id, person_id=actor.person_id,
                       size_bytes=len(payload), prepared=prepared)
        source.sha256 = digest
        source.created_at = db.scalar(text("SELECT transaction_timestamp()"))
        db.add(source)
        db.flush()
        source.status = "available"
        source.metadata_jsonb = {**source.metadata_jsonb, "completion": {
            "etag_sha256": "a" * 64, "head_manifest_sha256": "b" * 64,
            "verified_at": source.created_at.isoformat(),
        }}
        db.commit()
        key = sha256(uuid4().bytes).hexdigest()
        binding = {**defaults["import_binding_jsonb"], "source_file_id": str(source.id),
                   "source_sha256": digest, "count_key_sha256": _storage_hash(key)}
        job = FileJob(**{**defaults, "idempotency_key": key, "import_binding_jsonb": binding})
        db.add(job)
        db.commit()
        job_id = job.id
        job.status = "prevalidating"
        job.started_at = db.scalar(text("SELECT clock_timestamp()"))
        db.commit()
    authorized = AuthorizedOpeningCountSource(source.id, digest, len(payload), payload)
    arguments = dict(actor=actor, job_id=job_id, source=authorized, request_id="pg16-error-prepare")
    cases = []
    for label, value in (("formula", "=1+1"), ("control-character", "invalid\nmessage")):
        with Session(api) as db:
            def corrupt_preview(session, context, instances):
                for row in session.dirty:
                    if isinstance(row, FileJob) and row.import_preview_jsonb:
                        document = deepcopy(row.import_preview_jsonb)
                        document["errors"][0]["message"] = value
                        row.import_preview_jsonb = document
            event.listen(db, "before_flush", corrupt_preview)
            try:
                prepare_persisted_opening_count_error(db, **arguments)
                db.commit()
            except DBAPIError as error:
                assert error.orig.sqlstate == "23514"
                db.rollback()
            else:
                raise AssertionError("database accepted uncontrolled error text")
        cases.append("error-preview-" + label + "-refused")
    with Session(api) as db:
        first = prepare_persisted_opening_count_error(db, **arguments)
        assert first.new_write_allowed
        db.commit()
    with Session(api) as db:
        recovered = prepare_persisted_opening_count_error(db, **arguments)
        assert not recovered.new_write_allowed and recovered.artifact == first.artifact
        db.commit()
    artifact = first.artifact

    class PrivateObjects:
        provider_code = "aliyun_oss_v2"
        calls = []
        head = StoredObjectHead(artifact.storage_key, len(artifact.payload), MIME,
            {"sha256": artifact.sha256, "file-id": str(artifact.file_id)}, md5(artifact.payload).hexdigest())

        def put_opening_count_error(self, **kwargs):
            self.calls.append("put")
            return self.head

        def head_object(self, **kwargs):
            assert kwargs == {"storage_key": artifact.storage_key}
            self.calls.append("head")
            return self.head

    objects = PrivateObjects()
    write_or_recover_opening_count_error(objects, artifact=artifact, allow_new_put=first.new_write_allowed)
    head = write_or_recover_opening_count_error(objects, artifact=artifact, allow_new_put=recovered.new_write_allowed)
    assert objects.calls == ["put", "head"]
    cases.extend(("error-preparation-commits-once", "error-recovery-head-only"))
    for field, value in (("import_error_sha256", "f" * 64), ("import_preview_jsonb", {})):
        with Session(api) as db:
            try:
                setattr(db.get(FileJob, job_id), field, value)
                db.commit()
            except DBAPIError as error:
                assert error.orig.sqlstate == "23514"
                db.rollback()
            else:
                raise AssertionError("prepared error identity could be changed")
        cases.append("immutable-error-" + field)
    for user_id, expected in ((other_actor_id, "opening_import_job_not_found"),
                              (actor_id, "opening_import_actor_stale")):
        with Session(api) as db:
            reader = load_formal_principal(db, user_id)
            if user_id == actor_id:
                reader = replace(reader, authorization_version=reader.authorization_version + 1)
            try:
                publish_persisted_opening_count_error(db, actor=reader, job_id=job_id,
                    head=head, request_id="pg16-error-forbidden")
            except OpeningCountImportJobError as error:
                assert error.code == expected
                db.rollback()
            else:
                raise AssertionError("unauthorized error report publication")
        cases.append(expected)
    with Session(api) as db:
        try:
            publish_persisted_opening_count_error(db, actor=actor, job_id=job_id,
                head=replace(head, etag="0" * 32), request_id="pg16-error-wrong-head")
        except OpeningCountImportJobError as error:
            assert error.code == "opening_import_error_head_changed"
            db.rollback()
        else:
            raise AssertionError("wrong object accepted")
    cases.append("error-head-mismatch-refused")
    for label in ("hash", "size", "job-key"):
        with Session(api) as db:
            def corrupt_file(session, context, instances):
                for row in session.new:
                    if isinstance(row, FileObject):
                        if label == "hash":
                            row.sha256 = "f" * 64
                        elif label == "size":
                            row.size_bytes += 1
                        else:
                            row.metadata_jsonb = {**row.metadata_jsonb, "idempotency_key_hash": "f" * 64}
            event.listen(db, "before_flush", corrupt_file)
            try:
                publish_persisted_opening_count_error(db, actor=actor, job_id=job_id,
                    head=head, request_id="pg16-error-forged-file")
                db.commit()
            except DBAPIError as error:
                assert error.orig.sqlstate == "23514"
                assert "0141 error file must match its prepared import" in str(error.orig)
                db.rollback()
            else:
                raise AssertionError("database accepted error file detached from prepared job")
        cases.append("error-file-" + label + "-binding-refused")
    for commit in (False, True):
        with Session(api) as db:
            assert db.get(FileObject, artifact.file_id) is None
            file_id = publish_persisted_opening_count_error(db, actor=actor, job_id=job_id,
                head=head, request_id="pg16-error-publish")
            assert file_id == artifact.file_id
            if commit:
                db.commit()
            else:
                db.rollback()
        with Session(api) as db:
            job = db.get(FileJob, job_id)
            assert job.status == ("failed" if commit else "prevalidating")
            assert (db.get(FileObject, file_id) is not None) is commit
    cases.extend(("error-file-job-rollback-together", "error-file-job-commit-together"))
    with Session(api) as db:
        status = read_opening_count_import_job(db, actor=actor, job_id=job_id)
        assert status.status == "failed" and status.completion_id is None
        audits = db.scalars(select(AuditEvent).where(AuditEvent.aggregate_type == "file_job",
            AuditEvent.aggregate_id == str(job_id))).all()
        assert sorted(a.action for a in audits) == ["opening_count_import_error_prepared",
                                                  "opening_count_import_error_published"]
        try:
            confirm_persisted_opening_count_import(db, **arguments)
        except OpeningCountImportJobError as error:
            assert error.code == "opening_import_not_awaiting_confirmation"
        else:
            raise AssertionError("error job executed count")
    cases.extend(("error-audit-exactly-once", "error-job-cannot-confirm"))
    from pg16_opening_import_download_gate import exercise_import_error_download
    cases.extend(exercise_import_error_download(api, actor=actor, other_actor_id=other_actor_id,
                                               job_id=job_id, artifact=artifact))
    from app.formal_services.opening_count_import_intake import create_opening_count_import_job
    from app.opening_count_import_worker import process_one_opening_count_import
    from app.formal_services.file_storage import FileStorageError
    from uuid import UUID
    with Session(api) as db:
        queued_id, replayed = create_opening_count_import_job(db, actor=actor,
            source_file_id=source.id, task_id=UUID(binding["task_id"]),
            round_id=UUID(binding["round_id"]), scope_id=UUID(binding["scope_id"]),
            idempotency_key=uuid4().hex, request_id="pg16-error-worker-create")
        assert not replayed
        db.commit()

    class LostAcknowledgementStorage:
        provider_code = "aliyun_oss_v2"
        calls = []
        head = None

        def read_opening_count_source(self, **coordinates):
            assert coordinates["file_id"] == str(source.id) and coordinates["sha256"] == digest
            self.calls.append("source")
            return payload

        def put_opening_count_error(self, **coordinates):
            self.calls.append("put")
            self.head = StoredObjectHead(coordinates["storage_key"], len(coordinates["payload"]), MIME,
                {"sha256": coordinates["sha256"], "file-id": coordinates["file_id"]},
                md5(coordinates["payload"]).hexdigest())
            raise FileStorageError("synthetic lost PUT acknowledgement")

        def head_object(self, **coordinates):
            self.calls.append("head")
            assert self.head and coordinates["storage_key"] == self.head.storage_key
            return self.head

    transport = LostAcknowledgementStorage()
    try:
        process_one_opening_count_import(lambda: Session(api), storage=transport, job_id=queued_id)
    except FileStorageError:
        pass
    else:
        raise AssertionError("unknown PUT was reported as a completed import")
    with Session(api) as db:
        unknown = db.get(FileJob, queued_id)
        assert unknown.status == "prevalidating" and unknown.import_error_sha256
        assert unknown.error_file_id is None and unknown.import_completion_id is None
    recovered_worker = process_one_opening_count_import(lambda: Session(api), storage=transport, job_id=queued_id)
    assert recovered_worker.status == "failed" and recovered_worker.recovered
    assert transport.calls == ["source", "put", "source", "head"]
    terminal = process_one_opening_count_import(lambda: Session(api), storage=transport, job_id=queued_id)
    assert terminal.status == "failed" and terminal.recovered
    assert transport.calls == ["source", "put", "source", "head"]
    cases.extend(("worker-error-put-unknown-retained", "worker-error-recovery-never-reputs",
                  "worker-terminal-error-no-source-io"))
    from app.formal_services.opening_count_import_termination import cancel_opening_count_import

    def create_pending_error():
        original_key = uuid4().hex
        with Session(api) as db:
            identifier, _ = create_opening_count_import_job(db, actor=actor,
                source_file_id=source.id, task_id=UUID(binding["task_id"]),
                round_id=UUID(binding["round_id"]), scope_id=UUID(binding["scope_id"]),
                idempotency_key=original_key, request_id="pg16-prepared-terminal-create")
            db.commit()
        return identifier, original_key

    def cancel_pending(identifier, original_key):
        with Session(api) as db:
            result, replayed = cancel_opening_count_import(db, actor=actor, job_id=identifier,
                idempotency_key=original_key, request_id="pg16-prepared-cancel")
            db.commit()
            assert result.status == "cancelled" and not replayed

    for ending in ("cancel", "revocation"):
        identifier, original_key = create_pending_error()
        uncertain = LostAcknowledgementStorage()
        uncertain.calls = []
        try:
            process_one_opening_count_import(lambda: Session(api), storage=uncertain, job_id=identifier)
        except FileStorageError:
            pass
        else:
            raise AssertionError("synthetic unknown PUT did not retain its preparation")
        with Session(api) as db:
            prepared_job = db.get(FileJob, identifier)
            original = (prepared_job.import_error_sha256, prepared_job.import_error_size_bytes,
                        prepared_job.import_preview_jsonb, prepared_job.import_binding_jsonb)
            assert prepared_job.status == "prevalidating" and original[0]
        if ending == "revocation":
            with api.connect() as connection:
                transaction = connection.begin()
                connection.execute(text("UPDATE users SET authorization_version=authorization_version+1 WHERE id=:id"),
                                   {"id": actor.user_id})
                factory = lambda: Session(connection, join_transaction_mode="rollback_only")
                refused = process_one_opening_count_import(factory, storage=uncertain, job_id=identifier)
                assert refused.status == "failed"
                connection.exec_driver_sql("SET CONSTRAINTS ALL IMMEDIATE")
                with factory() as db:
                    failed_job = db.get(FileJob, identifier)
                    assert failed_job.error_detail == "opening_import_context_changed"
                    assert (failed_job.import_error_sha256, failed_job.import_error_size_bytes,
                            failed_job.import_preview_jsonb, failed_job.import_binding_jsonb) == original
                    assert failed_job.error_file_id is None and failed_job.import_completion_id is None
                transaction.rollback()
            cases.append("prepared-unknown-put-revocation-retains-evidence-no-io")
        cancel_pending(identifier, original_key)
        terminal = process_one_opening_count_import(lambda: Session(api), storage=uncertain, job_id=identifier)
        assert terminal.status == "cancelled" and uncertain.calls == ["source", "put"]
        with Session(api) as db:
            ended = db.get(FileJob, identifier)
            assert (ended.import_error_sha256, ended.import_error_size_bytes,
                    ended.import_preview_jsonb, ended.import_binding_jsonb) == original
            assert ended.error_file_id is None and ended.import_completion_id is None
        if ending == "cancel":
            cases.append("prepared-unknown-put-cancel-retains-evidence-no-io")

    concurrent_id, concurrent_key = create_pending_error()

    class CancelDuringPut(LostAcknowledgementStorage):
        def put_opening_count_error(self, **coordinates):
            try:
                super().put_opening_count_error(**coordinates)
            except FileStorageError:
                cancel_pending(concurrent_id, concurrent_key)
                return self.head

    concurrent = CancelDuringPut()
    concurrent.calls = []
    stopped = process_one_opening_count_import(lambda: Session(api), storage=concurrent, job_id=concurrent_id)
    assert stopped.status == "cancelled" and concurrent.calls == ["source", "put"]
    with Session(api) as db:
        ended = db.get(FileJob, concurrent_id)
        assert ended.error_file_id is None and ended.import_completion_id is None and ended.import_error_sha256
    cases.append("inflight-put-cancellation-never-publishes-or-counts")
    return cases
