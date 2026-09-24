from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
import pytest

from app.formal_services.opening_count_import_workbook import (
    HEADERS,
    SHEET_NAME,
    OpeningCountImportFormatError,
    prevalidate_opening_count_workbook,
    render_opening_count_error_report,
    render_opening_count_import_template,
)
from app.formal_services.opening_count_import_prevalidation import (
    _pending_verification_errors,
)


def _book(*rows):
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _row(**changes):
    values = {
        "物料标识": "ADQCPN0123", "标识类型": "sku_code", "成色": "new",
        "库存状态": "available", "实盘数量": "2.000", "批次号": None,
        "序列号": None, "序列标识类型": None, "差异原因": None, "备注": "现场实盘",
    }
    values.update(changes)
    return tuple(values[key] for key in HEADERS)


def test_valid_import_is_canonical_and_never_infers_zero_stock():
    data = _book(_row(), _row(**{"物料标识": "ADQCPN0456", "实盘数量": 1,
                               "序列号": "SN00001", "序列标识类型": "serial_no"}))
    result = prevalidate_opening_count_workbook(data)
    assert result.ready and result.row_count == 2 and not result.errors
    assert result.observation_source_rows == (2, 3)
    assert len(result.observations) == 2
    assert [item.counted_qty for item in result.observations] == [2, 1]
    assert all(item.count_method == "import" for item in result.observations)
    assert result.observations[1].serial_no_raw == "SN00001"
    assert len(result.source_sha256) == len(result.payload_sha256) == 64
    with pytest.raises(OpeningCountImportFormatError, match="没有盘点明细"):
        prevalidate_opening_count_workbook(render_opening_count_import_template())


def test_invalid_cells_produce_report_without_partial_executable_rows_or_source_echo():
    injected = "=HYPERLINK(\"https://example.invalid\",\"send\")"
    data = _book(_row(**{"物料标识": 12345, "实盘数量": "0", "备注": injected}),
                 _row(**{"物料标识": "ADQCPN2222", "序列号": "SN2"}))
    result = prevalidate_opening_count_workbook(data)
    assert not result.ready and result.observations == () and result.payload_sha256 is None
    assert {(item.row, item.field, item.code) for item in result.errors} >= {
        (2, "remark", "formula_forbidden"), (3, "row", "observation_invalid")
    }
    report = render_opening_count_error_report(result.errors)
    assert injected.encode() not in report and b"example.invalid" not in report
    book = load_workbook(BytesIO(report), read_only=True, data_only=False)
    try:
        rows = list(book.active.values)
        assert rows[0] == ("工作表", "行号", "字段", "错误代码", "说明")
        assert all(row[0] == SHEET_NAME for row in rows[1:])
        assert all(row[4] != injected for row in rows[1:])
    finally:
        book.close()


def test_blank_sheet_rows_preserve_actual_source_row_coordinates():
    result = prevalidate_opening_count_workbook(
        _book(_row(), (None,) * len(HEADERS), _row(**{"物料标识": "ADQCPN0456"}))
    )
    assert result.ready and result.row_count == 2
    assert result.observation_source_rows == (2, 4)


def test_business_pending_report_is_bounded_and_rejects_invalid_source_coordinates():
    source_rows = tuple(range(2, 1004))
    errors = _pending_verification_errors(source_rows, tuple(range(1, 1003)))
    assert len(errors) == 1000
    assert errors[0].row == 2
    assert errors[-1].code == "error_limit_reached"
    assert render_opening_count_error_report(errors)
    with pytest.raises(OpeningCountImportFormatError, match="行号与源文件不一致"):
        _pending_verification_errors(source_rows, (1003,))


@pytest.mark.parametrize("quantity", ["0", "-1", "1.2345", "NaN", "1e2", True])
def test_quantity_rejects_ambiguous_or_out_of_scale_values(quantity):
    result = prevalidate_opening_count_workbook(_book(_row(**{"实盘数量": quantity})))
    assert not result.ready and result.observations == ()
    assert any(item.field == "counted_qty" and item.code == "quantity_invalid" for item in result.errors)


def test_untrusted_structure_is_rejected_before_openpyxl_interprets_rows():
    with pytest.raises(OpeningCountImportFormatError, match="有效的 XLSX"):
        prevalidate_opening_count_workbook(b"not an xlsx")
    original = _book(_row())
    altered = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(altered, "w") as destination:
        for item in source.infolist():
            destination.writestr(item, source.read(item.filename))
        destination.writestr("../unexpected.xml", "x")
    with pytest.raises(OpeningCountImportFormatError, match="包路径无效"):
        prevalidate_opening_count_workbook(altered.getvalue())


def test_large_invalid_sheet_still_has_bounded_error_report():
    data = _book(*[_row(**{"物料标识": 12345, "实盘数量": "0"}) for _ in range(502)])
    result = prevalidate_opening_count_workbook(data)
    assert not result.ready and result.observations == ()
    assert len(result.errors) == 1000
    assert result.errors[-1].code == "error_limit_reached"
    assert render_opening_count_error_report(result.errors)


def test_extra_sheet_and_wrong_header_are_not_guessable():
    for mutation in (
        lambda book: book.create_sheet("历史导出"),
        lambda book: setattr(book.active["A1"], "value", "SKU"),
    ):
        book = Workbook()
        book.active.title = SHEET_NAME
        book.active.append(HEADERS)
        book.active.append(_row())
        mutation(book)
        output = BytesIO()
        book.save(output)
        book.close()
        with pytest.raises(OpeningCountImportFormatError):
            prevalidate_opening_count_workbook(output.getvalue())


def test_forged_short_worksheet_dimension_cannot_hide_observations():
    original = _book(_row(), _row(**{"物料标识": "ADQCPN0456"}))
    altered = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(altered, "w") as destination:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert b'ref="A1:J3"' in data
                data = data.replace(b'ref="A1:J3"', b'ref="A1:J2"', 1)
            destination.writestr(item, data)
    result = prevalidate_opening_count_workbook(altered.getvalue())
    assert result.ready and result.row_count == 2
    assert len(result.observations) == 2


def test_forged_short_worksheet_dimension_cannot_hide_extra_column():
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    sheet.append(_row())
    sheet["K2"] = "hidden"
    original = BytesIO()
    book.save(original)
    book.close()
    altered = BytesIO()
    with ZipFile(BytesIO(original.getvalue())) as source, ZipFile(altered, "w") as destination:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert b'ref="A1:K2"' in data
                data = data.replace(b'ref="A1:K2"', b'ref="A1:J2"', 1)
            destination.writestr(item, data)
    with pytest.raises(OpeningCountImportFormatError, match="列数或行数超出"):
        prevalidate_opening_count_workbook(altered.getvalue())
