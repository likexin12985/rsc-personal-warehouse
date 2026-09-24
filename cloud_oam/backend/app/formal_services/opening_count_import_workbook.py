"""Bounded XLSX format precheck for an opening-stocktake scope count.

This module does not authorize a scope, resolve OAM material identifiers, or
post stock. A future import job must retain the source hash and revalidate the
current scope and actor before it invokes the existing count command.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import re
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.read_only import EmptyCell
from pydantic import ValidationError

from ..opening_stocktake_schemas import OpeningPhysicalObservationIn
from .inventory_report_workbook import _canonicalize_package


SHEET_NAME = "期初盘点_V1"
HEADERS = (
    "物料标识", "标识类型", "成色", "库存状态", "实盘数量",
    "批次号", "序列号", "序列标识类型", "差异原因", "备注",
)
FIELDS = (
    "material_identifier_raw", "material_identifier_type", "condition_code",
    "availability_bucket", "counted_qty", "lot_no_raw", "serial_no_raw",
    "serial_identifier_type", "reason_code", "remark",
)
MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_ROWS = 10_000
MAX_ERRORS = 1_000
_MAX_QUANTITY = Decimal("1000000000000000")
_DECIMAL_TEXT = re.compile(r"^(0|[1-9]\d*)(?:\.\d{1,3})?$")


class OpeningCountImportFormatError(ValueError):
    """The workbook cannot be safely interpreted as this import format."""


@dataclass(frozen=True, slots=True)
class ImportRowError:
    row: int
    field: str
    code: str
    message: str


@dataclass(frozen=True, slots=True)
class OpeningCountImportPreview:
    source_sha256: str
    row_count: int
    observations: tuple[OpeningPhysicalObservationIn, ...]
    errors: tuple[ImportRowError, ...]
    payload_sha256: str | None
    observation_source_rows: tuple[int, ...] = ()

    @property
    def ready(self) -> bool:
        return (self.payload_sha256 is not None and not self.errors
                and len(self.observation_source_rows) == len(self.observations))


def prevalidate_opening_count_workbook(data: bytes) -> OpeningCountImportPreview:
    """Check file safety and row format, returning no executable rows on error."""

    if not isinstance(data, bytes) or not 0 < len(data) <= MAX_FILE_BYTES:
        raise OpeningCountImportFormatError("文件大小不在允许范围内")
    _validate_archive(data)
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise OpeningCountImportFormatError("无法读取 XLSX 文件") from exc
    try:
        if book.sheetnames != [SHEET_NAME] or book.active.sheet_state != "visible":
            raise OpeningCountImportFormatError("工作表名称或可见性不符合模板")
        sheet = book.active
        if (sheet.max_column is not None and sheet.max_column > len(HEADERS)
                or sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1):
            raise OpeningCountImportFormatError("列数或行数超出模板限制")
        # Read-only openpyxl otherwise trusts the untrusted <dimension ref>.
        # Recompute from actual worksheet rows so a forged short range cannot
        # hide extra observations or columns after the declared endpoint.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        first = next(iterator, ())
        if len(first) != len(HEADERS) or tuple(cell.value for cell in first) != HEADERS:
            raise OpeningCountImportFormatError("表头与期初盘点模板不一致")
        rows: list[OpeningPhysicalObservationIn] = []
        source_rows: list[int] = []
        errors: list[ImportRowError] = []
        row_count = 0
        for number, actual_cells in enumerate(iterator, start=2):
            if number > MAX_ROWS + 1 or len(actual_cells) > len(HEADERS):
                raise OpeningCountImportFormatError("列数或行数超出模板限制")
            cells = actual_cells + (EmptyCell(),) * (len(HEADERS) - len(actual_cells))
            if all(cell.value is None for cell in cells):
                continue
            row_count += 1
            for field, cell in zip(FIELDS, cells, strict=True):
                if cell.data_type == "f":
                    _add_error(errors, number, field, "formula_forbidden", "导入单元格不能使用公式")
            if any(cell.data_type == "f" for cell in cells):
                continue
            values: dict[str, object] = {"count_method": "import"}
            invalid = False
            for field, cell in zip(FIELDS, cells, strict=True):
                value = cell.value
                if field == "counted_qty":
                    try:
                        values[field] = _quantity_text(value)
                    except ValueError:
                        _add_error(errors, number, field, "quantity_invalid", "实盘数量须为大于零且最多三位小数")
                        invalid = True
                elif field in ("lot_no_raw", "serial_no_raw", "serial_identifier_type", "reason_code"):
                    if value is not None:
                        if not _valid_text(value):
                            _add_error(errors, number, field, "text_invalid", "该字段必须是无首尾空格的文本")
                            invalid = True
                        else:
                            values[field] = value
                elif field == "remark" and value is None:
                    values[field] = ""
                elif not _valid_text(value):
                    _add_error(errors, number, field, "text_invalid", "该字段必须是无首尾空格的文本")
                    invalid = True
                else:
                    values[field] = value
            if invalid:
                continue
            try:
                rows.append(OpeningPhysicalObservationIn.model_validate(values))
                source_rows.append(number)
            except ValidationError as exc:
                for issue in exc.errors():
                    field = str(issue["loc"][0]) if issue["loc"] else "row"
                    _add_error(errors, number, field, "observation_invalid", "盘点明细不符合字段或关联规则")
        if row_count == 0:
            raise OpeningCountImportFormatError("工作表没有盘点明细")
        source_hash = sha256(data).hexdigest()
        if errors:
            return OpeningCountImportPreview(source_hash, row_count, (), tuple(errors), None)
        payload = {"schema": "opening_count_import_v1", "observations": [
            row.model_dump(mode="json") for row in rows
        ]}
        canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
        return OpeningCountImportPreview(
            source_hash, row_count, tuple(rows), (), sha256(canonical).hexdigest(),
            tuple(source_rows),
        )
    except OpeningCountImportFormatError:
        raise
    except Exception as exc:
        raise OpeningCountImportFormatError("XLSX 工作表内容无法安全读取") from exc
    finally:
        book.close()


def render_opening_count_import_template() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        column[0].number_format = "@"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def render_opening_count_error_report(errors: tuple[ImportRowError, ...]) -> bytes:
    """Render stable controlled metadata for private, write-once error files."""

    if not errors or len(errors) > MAX_ERRORS:
        raise OpeningCountImportFormatError("错误报告行数无效")
    for item in errors:
        if (
            type(item) is not ImportRowError
            or type(item.row) is not int or not 2 <= item.row <= MAX_ROWS + 1
            or not isinstance(item.field, str) or item.field not in (*FIELDS, "row")
            or not isinstance(item.code, str)
            or not re.fullmatch(r"[a-z][a-z0-9_]{0,79}", item.code)
            or not isinstance(item.message, str) or not 1 <= len(item.message) <= 500
            or any(ord(char) < 32 for char in item.message)
            or item.message.lstrip().startswith(("=", "+", "-", "@"))
        ):
            raise OpeningCountImportFormatError("错误报告内容无效")
    book = Workbook()
    try:
        sheet = book.active
        sheet.title = "预校验错误"
        sheet.append(("工作表", "行号", "字段", "错误代码", "说明"))
        for item in errors:
            sheet.append((SHEET_NAME, item.row, item.field, item.code, item.message))
        output = BytesIO()
        book.save(output)
        # Regenerating after an unknown PUT must yield the same SHA-256 even
        # when wall time and ZIP timestamps changed. Never include raw cells.
        return _canonicalize_package(output.getvalue())
    finally:
        book.close()


def _validate_archive(data: bytes) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            names: set[str] = set()
            total = 0
            entries = archive.infolist()
            if not 0 < len(entries) <= 1_000:
                raise OpeningCountImportFormatError("XLSX 包条目数量无效")
            for item in entries:
                name = item.filename
                if (name in names or name.startswith("/") or "\\" in name
                    or any(part in ("", ".", "..") for part in name.split("/"))):
                    raise OpeningCountImportFormatError("XLSX 包路径无效")
                names.add(name)
                if item.flag_bits & 1 or item.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
                    raise OpeningCountImportFormatError("XLSX 包含不支持的加密或压缩条目")
                if name.lower().startswith("xl/externallinks/") or name.lower().endswith("vbaproject.bin"):
                    raise OpeningCountImportFormatError("XLSX 不能包含外部链接或宏")
                total += item.file_size
                if total > MAX_UNCOMPRESSED_BYTES or item.file_size > MAX_UNCOMPRESSED_BYTES:
                    raise OpeningCountImportFormatError("XLSX 解压后超过大小限制")
                if item.file_size and (not item.compress_size or item.file_size > item.compress_size * 250):
                    raise OpeningCountImportFormatError("XLSX 压缩比例异常")
            if "xl/workbook.xml" not in names or archive.testzip() is not None:
                raise OpeningCountImportFormatError("XLSX 包不完整")
    except OpeningCountImportFormatError:
        raise
    except Exception as exc:
        raise OpeningCountImportFormatError("文件不是有效的 XLSX") from exc


def _valid_text(value: object) -> bool:
    return (isinstance(value, str) and bool(value) and value == value.strip()
            and all(ord(char) >= 32 and ord(char) != 127 for char in value))


def _quantity_text(value: object) -> str:
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ValueError("invalid quantity")
    if isinstance(value, str) and _DECIMAL_TEXT.fullmatch(value) is None:
        raise ValueError("invalid quantity")
    try:
        quantity = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("invalid quantity") from exc
    if not quantity.is_finite() or quantity <= 0 or quantity >= _MAX_QUANTITY or quantity.as_tuple().exponent < -3:
        raise ValueError("invalid quantity")
    result = format(quantity, "f")
    return result.rstrip("0").rstrip(".") if "." in result else result


def _add_error(errors: list[ImportRowError], row: int, field: str, code: str, message: str) -> None:
    if len(errors) >= MAX_ERRORS:
        errors[-1] = ImportRowError(
            row, "row", "error_limit_reached", "仅保留前 999 项错误，其余请修正后重试"
        )
        return
    errors.append(ImportRowError(row, field, code, message))
